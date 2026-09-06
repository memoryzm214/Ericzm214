# -*- coding: utf-8 -*-
import os, re, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from text import XUANTI, JIANJIE, REFS, BUDGET, TOTAL
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SONG, HEI, KAI, TNR = '宋体', '黑体', '楷体', 'Times New Roman'
ACCENT = RGBColor(0x1F, 0x3A, 0x2E)
GREY = RGBColor(0x59, 0x5F, 0x5A)


def font(run, name=SONG, sz=12, bold=False, color=None):
    run.font.size = Pt(sz); run.font.bold = bold
    run.font.name = TNR
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    if color is not None:
        run.font.color.rgb = color


def cnt(blocks):
    return len(re.sub(r'\s', '', ''.join(t for k, t in blocks if k == 'p')))


def shade(cell, hexcolor):
    el = OxmlElement('w:shd'); el.set(qn('w:val'), 'clear'); el.set(qn('w:fill'), hexcolor)
    cell._tc.get_or_add_tcPr().append(el)


doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Cm(2.2)
sec.left_margin = sec.right_margin = Cm(2.6)
st = doc.styles['Normal']; st.font.name = TNR; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'), SONG)


def para(text, name=SONG, sz=12, bold=False, color=None, align=None,
         first=True, before=0, after=6, spacing=1.5):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    pf.line_spacing = spacing
    if first:
        pf.first_line_indent = Pt(sz * 2)
    if align is not None:
        p.alignment = align
    font(p.add_run(text), name, sz, bold, color)
    return p


def h1(num, title, note):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.space_before = Pt(18); pf.space_after = Pt(4); pf.line_spacing = 1.3
    font(p.add_run(f'{num}　{title}'), HEI, 14, True, ACCENT)
    q = doc.add_paragraph(); q.paragraph_format.space_after = Pt(8)
    q.paragraph_format.line_spacing = 1.3
    font(q.add_run(note), KAI, 10.5, False, GREY)


def h2(t):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.space_before = Pt(10); pf.space_after = Pt(3); pf.line_spacing = 1.4
    font(p.add_run(t), HEI, 12, True)


# ── 封面区 ──
para('湖北省社会科学基金一般项目', HEI, 16, True, ACCENT, WD_ALIGN_PARAGRAPH.CENTER, False, 0, 2, 1.2)
para('申报系统填报内容（拟稿）', HEI, 16, True, ACCENT, WD_ALIGN_PARAGRAPH.CENTER, False, 0, 10, 1.2)
para('课题名称：高校共享道路慢行主体的安全认知偏差与行为引导研究',
     KAI, 12, False, None, WD_ALIGN_PARAGRAPH.CENTER, False, 0, 2, 1.3)
para('申报年度：2026年度　　项目类别：一般项目　　资助额度：2万元',
     KAI, 11, False, GREY, WD_ALIGN_PARAGRAPH.CENTER, False, 0, 14, 1.3)

para('本稿依据《高校共享道路慢行主体的安全认知偏差与行为引导研究》研究报告（2026年9月4日版）拟定，'
     '所有数据、结论与文献均与报告正文一致。以下各节标题即申报系统中对应的填写字段，可直接复制粘贴。',
     KAI, 10.5, False, GREY, None, True, 0, 12, 1.4)

# ── 速查表 ──
h2('填报速查')
tb = doc.add_table(rows=5, cols=3); tb.style = 'Table Grid'; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = ['系统字段', '系统要求', '本稿字数／项数']
for j, t in enumerate(hdr):
    c = tb.rows[0].cells[j]; c.text = ''
    font(c.paragraphs[0].add_run(t), HEI, 10.5, True)
    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    shade(c, 'E8EBDE')
rows = [('（3）项目选题', '1500字左右；国内外研究现状述评、选题的意义', f'{cnt(XUANTI)} 字'),
        ('（4）项目内容简介', '2500字左右；基本思路和方法、主要观点和创新点', f'{cnt(JIANJIE)} 字'),
        ('（5）主要参考文献', '限填20项', f'{len(REFS)} 项'),
        ('6．经费信息', '单位为元，保留两位小数，总数2万元', f'{TOTAL} 元')]
for i, r in enumerate(rows, 1):
    for j, t in enumerate(r):
        c = tb.rows[i].cells[j]; c.text = ''
        p = c.paragraphs[0]; p.paragraph_format.line_spacing = 1.2
        font(p.add_run(t), SONG, 10.5)
        if j == 2:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for row in tb.rows:
    row.cells[0].width = Cm(3.6); row.cells[1].width = Cm(8.4); row.cells[2].width = Cm(3.4)

# ── 一、项目选题 ──
h1('一', '项目选题', f'【对应系统"项目信息—项目选题"字段。要点提示：1500字左右，'
                     f'课题国内外研究现状述评；选题的意义。本稿正文 {cnt(XUANTI)} 字。】')
for k, t in XUANTI:
    (h2(t) if k == 'h2' else para(t))

# ── 二、项目内容简介 ──
h1('二', '项目内容简介', f'【对应系统"项目信息—项目内容简介"字段。要点提示：2500字左右，'
                        f'课题研究的基本思路和方法，主要观点和创新点等。本稿正文 {cnt(JIANJIE)} 字。】')
for k, t in JIANJIE:
    (h2(t) if k == 'h2' else para(t))

# ── 三、主要参考文献 ──
h1('三', '主要参考文献', '【对应系统"项目信息—主要参考文献"字段。限填20项，本稿从研究报告的97条参考文献中'
                        '择要选取20条，覆盖共享空间理论、冲突识别技术、主观安全感知、行为决策理论、'
                        '风险认知、需求分类与可用性评估六个方面，其中中文文献5条、近五年文献11条。】')
for i, r in enumerate(REFS, 1):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.left_indent = Cm(0.9); pf.first_line_indent = Cm(-0.9)
    pf.space_after = Pt(3); pf.line_spacing = 1.35
    font(p.add_run(f'[{i}] {r}'), SONG, 10.5)

# ── 四、经费信息 ──
h1('四', '经费信息', '【对应系统"经费信息"页。单位：元，填写数字并保留两位小数；'
                    '各分项之和须等于合计 20000.00 元。系统只需填写"金额"一列，'
                    '"支出用途说明"供填报时自查与答辩备查之用。】')
bt = doc.add_table(rows=len(BUDGET) + 2, cols=3); bt.style = 'Table Grid'
bt.alignment = WD_TABLE_ALIGNMENT.CENTER
for j, t in enumerate(['经费科目', '金额（元）', '支出用途说明']):
    c = bt.rows[0].cells[j]; c.text = ''
    font(c.paragraphs[0].add_run(t), HEI, 10.5, True)
    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    shade(c, 'E8EBDE')
for i, (name, amt, use) in enumerate(BUDGET, 1):
    vals = [name, amt, use]
    for j, t in enumerate(vals):
        c = bt.rows[i].cells[j]; c.text = ''
        p = c.paragraphs[0]; p.paragraph_format.line_spacing = 1.2
        font(p.add_run(t), SONG, 10.5)
        if j == 1:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
last = bt.rows[len(BUDGET) + 1]
for j, t in enumerate(['合　计', TOTAL, '直接费用 18000.00 元，间接费用 2000.00 元']):
    c = last.cells[j]; c.text = ''
    p = c.paragraphs[0]; p.paragraph_format.line_spacing = 1.2
    font(p.add_run(t), HEI, 10.5, True)
    if j <= 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shade(c, 'F2F4EC')
for row in bt.rows:
    row.cells[0].width = Cm(4.6); row.cells[1].width = Cm(2.6); row.cells[2].width = Cm(8.2)

h2('预算编制说明')
para('本预算依据研究报告所载的实际工作量测算。劳务费占比最高（30.0%），'
     '与本课题以人为对象的数据采集规模相匹配：子任务①需132份问卷与53人次访谈，'
     '子任务②需232份有效问卷，子任务③需870条场景评分、30人次深度访谈与3场参与式工作坊，'
     '子任务④需86份Kano问卷与26人可用性测试，另有视频标注双人独立作业的助研劳务。'
     '专家咨询费用于两轮德尔菲专家咨询、量表内容效度评议与设计方案专家评审。'
     '设备费仅列视频采集配件的租赁与耗材，不购置大型设备。'
     '间接费用取2000.00元，占直接费用的11.1%，未超过省级社科基金的计提上限。')

# ── 五、图表 ──
h1('五', '可随申报材料提交的图表', '【系统的项目选题与内容简介两栏为纯文本输入，通常不支持插图；'
                                '以下图表建议在需要上传附件、活页或答辩汇报时使用，均可从研究报告中直接取用。】')
figs = [('图1  本课题技术路线图',
         '一页说清"特征识别—模型构建—需求识别—策略生成"四阶段与四项子任务的逐级供给关系，'
         '最适合放在活页或汇报首页。'),
        ('图5／图6  行人／骑行者行为决策模型',
         '两张同坐标同变量的路径图，直观呈现两类慢行主体主导传导链与解释力的差异，'
         '用于支撑"差异化引导"这一核心论点。'),
        ('图8  安全引导系统需求清单（简版）',
         '三类需求自下而上分层，右侧箭头标注"先基础后提升"的实施顺序，'
         '适合说明成果的可落地性与分步实施建议。'),
        ('表6  本课题研究设计矩阵',
         '四项子任务的研究问题、主要方法、数据来源与主要产出一览，'
         '可作为内容简介的补充材料，便于评审快速核对工作量。')]
for n, d in figs:
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.left_indent = Cm(0.6); pf.space_after = Pt(2); pf.line_spacing = 1.4
    font(p.add_run('● ' + n + '　'), HEI, 11, True)
    font(p.add_run(d), SONG, 11)

# ── 六、注意事项 ──
h1('六', '填报注意事项', '【系统录入前请逐条自查。】')
tips = ['本稿的引号、破折号与省略号均为中文全角符号，从Word复制到系统输入框时'
        '可能被转为半角或出现乱码，粘贴后请通读一遍；建议先粘贴到记事本转为纯文本，再粘入系统。',
        '项目选题与内容简介两栏若系统设有硬性字符上限，可优先压缩"国内外研究现状述评"'
        '第（一）至（四）小节中的方法罗列部分，保留三处研究缺口与选题意义。',
        '主要参考文献限填20项，本稿已选足20条；若系统按行分栏输入，请逐条录入，不要合并。',
        '经费各栏必须填写数字并保留两位小数，未发生的科目填 0.00，不可留空；'
        '录入后核对合计栏是否自动显示为 20000.00。',
        '课题名称、起止时间、成果形式等基本信息须与纸质申报书完全一致，'
        '提交前建议逐栏与申报书对照一次。']
for i, t in enumerate(tips, 1):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.left_indent = Cm(0.9); pf.first_line_indent = Cm(-0.9)
    pf.space_after = Pt(4); pf.line_spacing = 1.45
    font(p.add_run(f'{i}．{t}'), SONG, 11)

out = os.path.join(os.path.dirname(os.path.abspath(__file__)), '申报系统填报内容_2026.docx')
doc.save(out)
print('saved', out)
