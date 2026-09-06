# -*- coding: utf-8 -*-
"""输出规范化修订版参考文献（97条）。改动过的条目加浅黄底纹，便于比对。"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_LINE_SPACING, WD_COLOR_INDEX
from docx.oxml.ns import qn

REFS = {
1:"中共中央, 国务院. 教育强国建设规划纲要（2024—2035年）[Z]. 北京: 人民出版社, 2025.",
2:"武汉市人民政府. 武汉市非机动车管理办法: 武汉市人民政府令第293号（2024年第三次修改）[Z]. 武汉: 武汉市人民政府, 2024.",
3:"全国人民代表大会常务委员会. 中华人民共和国道路交通安全法: 2021年修正[Z]. 北京: 中国法制出版社, 2021.",
4:"国家市场监督管理总局, 国家标准化管理委员会. 电动自行车安全技术规范: GB 17761—2024[S]. 北京: 中国标准出版社, 2024.",
5:"中华人民共和国住房和城乡建设部. 城市步行和自行车交通系统规划标准: GB/T 51439—2021[S]. 北京: 中国建筑工业出版社, 2021.",
6:"中华人民共和国住房和城乡建设部. 城市综合交通体系规划标准: GB/T 51328—2018[S]. 北京: 中国建筑工业出版社, 2018.",
7:"HAMILTON-BAILLIE B. Shared space: reconciling people, places and traffic[J]. Built Environment, 2008, 34(2): 161-181. DOI: 10.2148/benv.34.2.161.",
8:"MOODY S, MELIA S. Shared space: research, policy and problems[J]. Proceedings of the Institution of Civil Engineers - Transport, 2014, 167(6): 384-392. DOI: 10.1680/tran.12.00047.",
9:"MOHER D, LIBERATI A, TETZLAFF J, et al. Preferred reporting items for systematic reviews and meta-analyses: the PRISMA statement[J]. BMJ, 2009, 339: b2535. DOI: 10.1136/bmj.b2535.",
10:"PAGE M J, MCKENZIE J E, BOSSUYT P M, et al. The PRISMA 2020 statement: an updated guideline for reporting systematic reviews[J]. BMJ, 2021, 372: n71. DOI: 10.1136/bmj.n71.",
11:"KARNDACHARUK A, WILSON D J, DUNN R. A review of the evolution of shared (street) space concepts in urban environments[J]. Transport Reviews, 2014, 34(2): 190-220. DOI: 10.1080/01441647.2014.893038.",
12:"中华人民共和国住房和城乡建设部. 无障碍设计规范: GB 50763—2012[S]. 北京: 中国建筑工业出版社, 2012.",
13:"中华人民共和国住房和城乡建设部. 建筑与市政工程无障碍通用规范: GB 55019—2021[S]. 北京: 中国建筑工业出版社, 2021.",
14:"PELTZMAN S. The effects of automobile safety regulation[J]. Journal of Political Economy, 1975, 83(4): 677-725. DOI: 10.1086/260352.",
15:"WILDE G J S. The theory of risk homeostasis: implications for safety and health[J]. Risk Analysis, 1982, 2(4): 209-225. DOI: 10.1111/j.1539-6924.1982.tb01384.x.",
16:"HYDÉN C. The development of a method for traffic safety evaluation: the Swedish traffic conflicts technique: Bulletin 70[R]. Lund: Lund University, 1987.",
17:"HAYWARD J C. Near-miss determination through use of a scale of danger[J]. Highway Research Record, 1972, 384: 24-34.",
18:"ALLEN B L, SHIN B T, COOPER P J. Analysis of traffic conflicts and collisions[J]. Transportation Research Record, 1978, 667: 67-74.",
19:"GETTMAN D, HEAD L. Surrogate safety measures from traffic simulation models[J]. Transportation Research Record, 2003, 1840(1): 104-115. DOI: 10.3141/1840-12.",
20:"LAURESHYN A, SVENSSON A, HYDEN C. Evaluation of traffic safety, based on micro-level behavioural data: theoretical framework and first implementation[J]. Accident Analysis & Prevention, 2010, 42(6): 1637-1646. DOI: 10.1016/j.aap.2010.03.021.",
21:"JOHNSSON C, LAURESHYN A, DE CEUNYNCK T. In search of surrogate safety indicators for vulnerable road users: a review of surrogate safety indicators[J]. Transport Reviews, 2018, 38(6): 765-785. DOI: 10.1080/01441647.2018.1442888.",
22:"ZHENG L, ISMAIL K, MENG X. Traffic conflict techniques for road safety analysis: open questions and some insights[J]. Canadian Journal of Civil Engineering, 2014, 41(7): 633-641. DOI: 10.1139/cjce-2013-0558.",
23:"MIAO L, LIU F, DENG Y. Analysis of traffic conflicts on slow-moving shared paths in Shenzhen, China[J]. Sustainability, 2025, 17(9): 4095. DOI: 10.3390/su17094095.",
24:"郑玉冰, 马羊, 程建川, 等. 基于轨迹数据的非机动车道内冲突事件自动识别与可视化[J]. 中国公路学报, 2022, 35(1): 71-84. DOI: 10.19721/j.cnki.1001-7372.2022.01.007.",
25:"CERVERO R, KOCKELMAN K. Travel demand and the 3Ds: density, diversity, and design[J]. Transportation Research Part D: Transport and Environment, 1997, 2(3): 199-219. DOI: 10.1016/s1361-9209(97)00009-6.",
26:"EWING R, HANDY S. Measuring the unmeasurable: urban design qualities related to walkability[J]. Journal of Urban Design, 2009, 14(1): 65-84. DOI: 10.1080/13574800802451155.",
27:"FRANK L D, SALLIS J F, SAELENS B E, et al. The development of a walkability index: application to the Neighborhood Quality of Life Study[J]. British Journal of Sports Medicine, 2010, 44(13): 924-933. DOI: 10.1136/bjsm.2009.058701.",
28:"ZHANG Z, FISHER T, WANG H. Walk Score, environmental quality and walking in a campus setting[J]. Land, 2023, 12(4): 732. DOI: 10.3390/land12040732.",
29:"KELLSTEDT D K, SPENGLER J O, MADDOCK J E. Comparing perceived and objective measures of bikeability on a university campus: a case study[J]. SAGE Open, 2021, 11(2): 21582440211018685. DOI: 10.1177/21582440211018685.",
30:"李琳, 叶宇, 陈泳. 城市步行安全及其环境影响要素研究综述与展望[J]. 风景园林, 2025, 32(2): 86-94. DOI: 10.3724/j.fjyl.202405210283.",
31:"李聪颖, 张洪涛, 李坤, 等. 城市自行车交通系统出行品质评价方法综述[J]. 交通运输工程学报, 2024, 24(6): 43-65. DOI: 10.19818/j.cnki.1671-1637.2024.06.003.",
32:"LANDIS B W, VATTIKUTI V R, BRANNICK M T. Real-time human perceptions: toward a bicycle level of service[J]. Transportation Research Record, 1997, 1578(1): 119-126. DOI: 10.3141/1578-15.",
33:"SORTON A, WALSH T. Bicycle stress level as a tool to evaluate urban and suburban bicycle compatibility[J]. Transportation Research Record, 1994, 1438: 17-24.",
34:"MEKURIA M C, FURTH P G, NIXON H. Low-stress bicycling and network connectivity: MTI Report 11-19[R]. San Jose: Mineta Transportation Institute, 2012.",
35:"HASSANPOUR A, BIGAZZI A. Clustering micromobility devices based on speed and comfort[J]. Findings, 2024: 123208. DOI: 10.32866/001c.123208.",
36:"金俊, 林思铭, 周珏, 等. “环境—感知”一体化的高强度商务区步行空间设计研究——长三角城市的实证探索[J]. 中国园林, 2025, 41(12): 41-48. DOI: 10.19775/j.cla.2025.12.0041.",
37:"朱萌, 王灿祥, 陈锦富. 街道步行环境夜间安全感关键影响因素识别——基于可穿戴生理传感器的感知实验[J]. 中国园林, 2023, 39(6): 64-69.",
38:"FISHBEIN M, AJZEN I. Belief, attitude, intention and behavior: an introduction to theory and research[M]. Reading, MA: Addison-Wesley, 1975.",
39:"AJZEN I. The theory of planned behavior[J]. Organizational Behavior and Human Decision Processes, 1991, 50(2): 179-211. DOI: 10.1016/0749-5978(91)90020-T.",
40:"BANDURA A. Self-efficacy: toward a unifying theory of behavioral change[J]. Psychological Review, 1977, 84(2): 191-215. DOI: 10.1037/0033-295x.84.2.191.",
41:"PARKER D, MANSTEAD A S R, STRADLING S G, et al. Intention to commit driving violations: an application of the theory of planned behavior[J]. Journal of Applied Psychology, 1992, 77(1): 94-101. DOI: 10.1037/0021-9010.77.1.94.",
42:"REASON J, MANSTEAD A, STRADLING S, et al. Errors and violations on the roads: a real distinction?[J]. Ergonomics, 1990, 33(10-11): 1315-1332. DOI: 10.1080/00140139008925335.",
43:"ROSENSTOCK I M. Historical origins of the health belief model[J]. Health Education Monographs, 1974, 2(4): 328-335. DOI: 10.1177/109019817400200403.",
44:"ROGERS R W. A protection motivation theory of fear appeals and attitude change[J]. Journal of Psychology, 1975, 91(1): 93-114. DOI: 10.1080/00223980.1975.9915803.",
45:"DAVIS F D. Perceived usefulness, perceived ease of use, and user acceptance of information technology[J]. MIS Quarterly, 1989, 13(3): 319-340. DOI: 10.2307/249008.",
46:"VENKATESH V, DAVIS F D. A theoretical extension of the technology acceptance model: four longitudinal field studies[J]. Management Science, 2000, 46(2): 186-204. DOI: 10.1287/mnsc.46.2.186.11926.",
47:"VENKATESH V, MORRIS M G, DAVIS G B, et al. User acceptance of information technology: toward a unified view[J]. MIS Quarterly, 2003, 27(3): 425-478. DOI: 10.2307/30036540.",
48:"CHEN H, GUO Y, LI L. Promoting sustainable mobility on campus: uncovering the behavioral mechanisms behind non-compliant e-bike use among university students[J]. Sustainability, 2025, 17(15): 7147. DOI: 10.3390/su17157147.",
49:"裴玉龙, 龙钰, 马丹. 交通安全意识对非机动车骑行者危险骑行行为的影响研究[J]. 交通信息与安全, 2024, 42(1): 49-58. DOI: 10.3963/j.jssn.1674-4861.2024.01.006.",
50:"VON STÜLPNAGEL R, RINTELEN H. A matter of space and perspective: cyclists', car drivers' and pedestrians' assumptions about subjective safety in shared traffic situations[J]. Transportation Research Part A: Policy and Practice, 2024, 179: 103941. DOI: 10.1016/j.tra.2023.103941.",
51:"刘孟歆, 秦华, 岳晨, 等. 右转车辆与过街行人交互过程的影响因素研究[J]. 包装工程, 2023, 44(12): 118-125. DOI: 10.19554/j.cnki.1001-3563.2023.12.012.",
52:"SLOVIC P. Perception of risk[J]. Science, 1987, 236(4799): 280-285. DOI: 10.1126/science.3563507.",
53:"TVERSKY A, KAHNEMAN D. Judgment under uncertainty: heuristics and biases[J]. Science, 1974, 185(4157): 1124-1131. DOI: 10.1126/science.185.4157.1124.",
54:"吕能超, 王玉刚, 周颖, 等. 道路交通安全分析与评价方法综述[J]. 中国公路学报, 2023, 36(4): 183-201. DOI: 10.19721/j.cnki.1001-7372.2023.04.016.",
55:"陈雅楠, 赵晓华, 李佳, 等. 基于科学知识图谱的道路交叉口安全设施设计综述及范式研究[J]. 北京工业大学学报, 2024, 50(12): 1501-1520. DOI: 10.11936/bjutxb2023030017.",
56:"LYNCH K. The image of the city[M]. Cambridge, MA: MIT Press, 1960.",
57:"PASSINI R. Wayfinding in architecture[M]. New York: Van Nostrand Reinhold, 1984.",
58:"ARTHUR P, PASSINI R. Wayfinding: people, signs, and architecture[M]. New York: McGraw-Hill, 1992.",
59:"MILLER G A. The magical number seven, plus or minus two: some limits on our capacity for processing information[J]. Psychological Review, 1956, 63(2): 81-97. DOI: 10.1037/h0043158.",
60:"SWELLER J. Cognitive load during problem solving: effects on learning[J]. Cognitive Science, 1988, 12(2): 257-285. DOI: 10.1207/s15516709cog1202_4.",
61:"WICKENS C D. Multiple resources and performance prediction[J]. Theoretical Issues in Ergonomics Science, 2002, 3(2): 159-177. DOI: 10.1080/14639220210123806.",
62:"WARE C. Information visualization: perception for design[M]. 3rd ed. Waltham: Morgan Kaufmann, 2013.",
63:"ELLIOT A J, MAIER M A. Color psychology: effects of perceiving color on psychological functioning in humans[J]. Annual Review of Psychology, 2014, 65: 95-120. DOI: 10.1146/annurev-psych-010213-115035.",
64:"THALER R H, SUNSTEIN C R. Nudge: improving decisions about health, wealth, and happiness[M]. New Haven: Yale University Press, 2008.",
65:"JEFFERY C R. Crime prevention through environmental design[M]. Beverly Hills: Sage, 1971.",
66:"NEWMAN O. Defensible space: crime prevention through urban design[M]. New York: Macmillan, 1972.",
67:"NORMAN D A, DRAPER S W. User centered system design: new perspectives on human-computer interaction[M]. Hillsdale, NJ: Lawrence Erlbaum, 1986.",
68:"SANDERS E B N, STAPPERS P J. Co-creation and the new landscapes of design[J]. CoDesign, 2008, 4(1): 5-18. DOI: 10.1080/15710880701875068.",
69:"NIELSEN J. Usability engineering[M]. Boston: Academic Press, 1993.",
70:"GIBSON J J. The ecological approach to visual perception[M]. Boston: Houghton Mifflin, 1979.",
71:"张乐, 汤晓敏. 可供性视角下大学老校区公共空间评价与更新设计指引——以上海交通大学徐汇校区为例[J]. 中国园林, 2025, 41(5): 115-122.",
72:"DALKEY N, HELMER O. An experimental application of the Delphi method to the use of experts[J]. Management Science, 1963, 9(3): 458-467. DOI: 10.1287/mnsc.9.3.458.",
73:"BRAUN V, CLARKE V. Using thematic analysis in psychology[J]. Qualitative Research in Psychology, 2006, 3(2): 77-101. DOI: 10.1191/1478088706qp063oa.",
74:"STRAUSS A, CORBIN J. Basics of qualitative research: techniques and procedures for developing grounded theory[M]. 2nd ed. Thousand Oaks: Sage, 1998.",
75:"COCHRAN W G. Sampling techniques[M]. 3rd ed. New York: John Wiley & Sons, 1977.",
76:"HU L, BENTLER P M. Cutoff criteria for fit indexes in covariance structure analysis: conventional criteria versus new alternatives[J]. Structural Equation Modeling, 1999, 6(1): 1-55. DOI: 10.1080/10705519909540118.",
77:"NUNNALLY J C, BERNSTEIN I H. Psychometric theory[M]. 3rd ed. New York: McGraw-Hill, 1994.",
78:"FORNELL C, LARCKER D F. Evaluating structural equation models with unobservable variables and measurement error[J]. Journal of Marketing Research, 1981, 18(1): 39-50. DOI: 10.1177/002224378101800104.",
79:"朱海腾, 李川云. 共同方法变异是“致命瘟疫”吗？——论争、新知与应对[J]. 心理科学进展, 2019, 27(4): 587-599. DOI: 10.3724/SP.J.1042.2019.00587.",
80:"HENSELER J, RINGLE C M, SARSTEDT M. A new criterion for assessing discriminant validity in variance-based structural equation modeling[J]. Journal of the Academy of Marketing Science, 2015, 43(1): 115-135. DOI: 10.1007/s11747-014-0403-8.",
81:"PREACHER K J, HAYES A F. Asymptotic and resampling strategies for assessing and comparing indirect effects in multiple mediator models[J]. Behavior Research Methods, 2008, 40(3): 879-891. DOI: 10.3758/brm.40.3.879.",
82:"温忠麟, 叶宝娟. 中介效应分析：方法和模型发展[J]. 心理科学进展, 2014, 22(5): 731-745. DOI: 10.3724/SP.J.1042.2014.00731.",
83:"温忠麟, 方杰, 谢晋艳, 等. 国内中介效应的方法学研究[J]. 心理科学进展, 2022, 30(8): 1692-1702. DOI: 10.3724/SP.J.1042.2022.01692.",
84:"CHEUNG G W, RENSVOLD R B. Evaluating goodness-of-fit indexes for testing measurement invariance[J]. Structural Equation Modeling, 2002, 9(2): 233-255. DOI: 10.1207/s15328007sem0902_5.",
85:"中华人民共和国住房和城乡建设部. 城市道路工程设计规范: CJJ 37—2012: 2016年版[S]. 北京: 中国建筑工业出版社, 2016.",
86:"COHEN J. Statistical power analysis for the behavioral sciences[M]. 2nd ed. Hillsdale, NJ: Lawrence Erlbaum, 1988.",
87:"MARTILLA J A, JAMES J C. Importance-performance analysis[J]. Journal of Marketing, 1977, 41(1): 77-79. DOI: 10.1177/002224297704100112.",
88:"KANO N, SERAKU N, TAKAHASHI F, et al. Attractive quality and must-be quality[J]. Journal of the Japanese Society for Quality Control, 1984, 14(2): 147-156. DOI: 10.20684/quality.14.2_147.",
89:"Design Council. Eleven lessons: managing design in eleven global brands——a study of the design process[R]. London: Design Council, 2007.",
90:"NORMAN D A. The design of everyday things[M]. Revised and expanded ed. New York: Basic Books, 2013.",
91:"中华人民共和国住房和城乡建设部. 建筑照明设计标准: GB/T 50034—2024[S]. 北京: 中国建筑工业出版社, 2024.",
92:"中华人民共和国住房和城乡建设部. 城市道路照明设计标准: CJJ 45—2015[S]. 北京: 中国建筑工业出版社, 2015.",
93:"国家市场监督管理总局, 国家标准化管理委员会. 道路交通标志和标线 第3部分: 道路交通标线: GB 5768.3—2025[S]. 北京: 中国标准出版社, 2025.",
94:"BROOKE J. SUS: a “quick and dirty” usability scale[M]//JORDAN P W, THOMAS B, WEERDMEESTER B A, et al. Usability evaluation in industry. London: Taylor & Francis, 1996: 189-194. DOI: 10.1201/9781498710411-35.",
95:"BANGOR A, KORTUM P T, MILLER J T. An empirical evaluation of the System Usability Scale[J]. International Journal of Human-Computer Interaction, 2008, 24(6): 574-594. DOI: 10.1080/10447310802205776.",
96:"BANGOR A, KORTUM P T, MILLER J T. Determining what individual SUS scores mean: adding an adjective rating scale[J]. Journal of Usability Studies, 2009, 4(3): 114-123.",
97:"BROOKE J. SUS: a retrospective[J]. Journal of Usability Studies, 2013, 8(2): 29-40.",
}
CHANGED = {1,2,3,4,16,30,34,49,51,54,55,69,85,89,91,93,96}

CN, EN = '宋体', 'Times New Roman'


def setfont(run, sz=10.5, bold=False):
    run.font.size = Pt(sz)
    run.font.bold = bold
    run.font.name = EN
    run._element.rPr.rFonts.set(qn('w:eastAsia'), CN)


doc = Document()
st = doc.styles['Normal']
st.font.name = EN
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), CN)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(6)
setfont(p.add_run('参考文献（规范化修订版·97条）'), 15, True)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(14)
r = p.add_run('说明：黄色底纹的 17 条为本次修改过著录格式或补充 DOI 的条目，'
              '核对无误后请在 Word 中全选并清除突出显示（开始→文本突出显示颜色→无颜色）。'
              '插入正文前，务必先按本表重排正文中的方括号引注编号。')
setfont(r, 9)
r.font.color.rgb = RGBColor(0x8E, 0x42, 0x38)

for i in range(1, 98):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.95)
    pf.first_line_indent = Cm(-0.95)
    pf.space_after = Pt(3)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    r = p.add_run(f'[{i}] {REFS[i]}')
    setfont(r)
    if i in CHANGED:
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW

out = os.path.join(os.path.dirname(__file__), '参考文献_规范化修订版.docx')
doc.save(out)
print('saved', out, len(REFS))
